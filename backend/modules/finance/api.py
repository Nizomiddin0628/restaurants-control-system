"""Moliya va hisobotlar API — /api/v1/finance/..."""
from __future__ import annotations

from datetime import date, timedelta
from io import BytesIO
from typing import Optional

from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from ninja import File, Router, Schema
from ninja.errors import HttpError
from ninja.files import UploadedFile

from core.audit import record
from core.auth import auth, require_module, require_perm

from . import reports
from .models import Expense, ExpenseCategory, ensure_categories

router = Router(tags=["finance"])


def _guard(request, perm: str):
    require_module(request, "finance")
    require_perm(request, perm)


def _period(start: Optional[date], end: Optional[date]) -> tuple[date, date]:
    end = end or timezone.localdate()
    start = start or (end - timedelta(days=29))
    if start > end:
        raise HttpError(400, "Boshlanish sanasi tugashdan keyin.")
    return start, end


class ExpenseIn(Schema):
    date: Optional[date] = None
    branch_id: Optional[int] = None
    category_id: int
    amount: int
    note: str = ""


class ExpenseOut(Schema):
    id: int
    date: date
    branch_id: Optional[int] = None
    branch_name: Optional[str] = None
    category_id: int
    category_name: str = ""
    amount: int
    note: str
    receipt: Optional[str] = None
    source: str
    ref: str
    created_by: Optional[str] = None

    @staticmethod
    def resolve_branch_name(obj):
        return obj.branch.name if obj.branch_id else None

    @staticmethod
    def resolve_category_name(obj):
        return obj.category.name

    @staticmethod
    def resolve_receipt(obj):
        return obj.receipt.url if obj.receipt else None

    @staticmethod
    def resolve_created_by(obj):
        return obj.created_by.full_name if obj.created_by_id else None


class CategoryIn(Schema):
    name: str
    is_fixed: bool = False
    kind: str = "opex"


# ------------------------------------------------------------------ hisobotlar
@router.get("/pnl", auth=auth)
def pnl(request, start: Optional[date] = None, end: Optional[date] = None, branch_id: Optional[int] = None):
    """Foyda-zarar: daromad − tannarx − mehnat − chiqimlar. Food cost %, prime cost %, sof marja."""
    _guard(request, "finance.view")
    s, e = _period(start, end)
    return reports.pnl(s, e, branch_id)


@router.get("/dashboard", auth=auth)
def dashboard(request, start: Optional[date] = None, end: Optional[date] = None, branch_id: Optional[int] = None):
    """Hisobotlar sahifasi uchun hammasi bir so'rovda."""
    _guard(request, "finance.view")
    s, e = _period(start, end)
    days = (e - s).days + 1
    prev_s, prev_e = s - timedelta(days=days), s - timedelta(days=1)
    cur, prev = reports.pnl(s, e, branch_id), reports.pnl(prev_s, prev_e, branch_id)
    def delta(k):  # oldingi davr bo'sh yoki juda kichik bo'lsa foiz ma'nosiz — None
        if not prev.get(k) or prev[k] < 0:
            return None
        v = round(100 * (cur[k] - prev[k]) / prev[k], 1)
        return v if abs(v) <= 500 else None
    return {
        "pnl": cur, "prev": {"revenue": prev["revenue"], "net_profit": prev["net_profit"], "orders": prev["orders"]},
        "delta": {"revenue": delta("revenue"), "net_profit": delta("net_profit"), "orders": delta("orders"), "avg_check": delta("avg_check")},
        "daily": reports.daily_series(s, e, branch_id),
        "hourly": reports.hourly_profile(s, e, branch_id),
        "top_products": reports.top_products(s, e, branch_id),
        "by_method": reports.by_payment_method(s, e, branch_id),
        "by_branch": reports.by_branch(s, e),
        "menu_engineering": reports.menu_engineering(s, e, branch_id),
    }


@router.get("/export.xlsx", auth=auth)
def export_xlsx(request, start: Optional[date] = None, end: Optional[date] = None, branch_id: Optional[int] = None):
    """Excel: P&L + kunlik savdo + top taomlar + chiqimlar — buxgalter/investor uchun."""
    _guard(request, "finance.view")
    from openpyxl import Workbook
    from openpyxl.styles import Font

    s, e = _period(start, end)
    p = reports.pnl(s, e, branch_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "P&L"
    bold = Font(bold=True)
    ws.append([f"Foyda-zarar hisoboti {s} — {e}"]); ws["A1"].font = bold
    for label, key, pkey in [("Daromad", "revenue", None), ("Buyurtmalar", "orders", None), ("O'rtacha chek", "avg_check", None),
                             ("Tannarx (COGS)", "cogs", "food_cost_percent"), ("Yalpi foyda", "gross_profit", "gross_margin_percent"),
                             ("Mehnat (oylik)", "labor", "labor_percent"), ("Chiqimlar", "expenses_total", "expenses_percent"),
                             ("Prime cost", "prime_cost", "prime_cost_percent"), ("Sof foyda", "net_profit", "net_margin_percent")]:
        ws.append([label, p[key], f"{p[pkey]}%" if pkey else ""])
    ws.append([]); ws.append(["Chiqimlar bo'yicha"]); ws[f"A{ws.max_row}"].font = bold
    for x in p["expenses"]:
        ws.append([x["category"], x["amount"]])
    ws2 = wb.create_sheet("Kunlik savdo"); ws2.append(["Sana", "Daromad", "Tannarx", "Buyurtmalar"])
    for d in reports.daily_series(s, e, branch_id):
        ws2.append([d["date"], d["revenue"], d["cogs"], d["orders"]])
    ws3 = wb.create_sheet("Top taomlar"); ws3.append(["Taom", "Soni", "Daromad", "Tannarx", "Marja", "Ulush %", "Sinf"])
    for t in reports.menu_engineering(s, e, branch_id):
        ws3.append([t["name"], t["qty"], t["revenue"], t["cost"], t["margin"], t["share"], t["class"]])
    ws4 = wb.create_sheet("Chiqimlar"); ws4.append(["Sana", "Kategoriya", "Summa", "Izoh", "Manba"])
    for x in Expense.objects.filter(date__gte=s, date__lte=e).select_related("category"):
        ws4.append([x.date, x.category.name, x.amount, x.note, x.source])
    for w in wb.worksheets:
        for col in w.columns:
            w.column_dimensions[col[0].column_letter].width = 22
    buf = BytesIO(); wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="hisobot_{s}_{e}.xlsx"'
    return resp


# ------------------------------------------------------------------ chiqimlar
@router.get("/categories", auth=auth)
def categories(request):
    _guard(request, "finance.view")
    ensure_categories()
    return [{"id": c.id, "code": c.code, "name": c.name, "is_fixed": c.is_fixed, "kind": c.kind} for c in ExpenseCategory.objects.all()]


@router.post("/categories", auth=auth)
def create_category(request, data: CategoryIn):
    _guard(request, "finance.edit")
    base = "".join(ch if ch.isalnum() else "_" for ch in data.name.lower())[:30] or "cat"
    code, i = base, 1
    while ExpenseCategory.objects.filter(code=code).exists():
        i += 1; code = f"{base}_{i}"
    c = ExpenseCategory.objects.create(code=code, name=data.name, is_fixed=data.is_fixed, kind=data.kind, sort_order=ExpenseCategory.objects.count())
    return {"id": c.id, "code": c.code, "name": c.name, "is_fixed": c.is_fixed, "kind": c.kind}


@router.get("/expenses", response=list[ExpenseOut], auth=auth)
def list_expenses(request, start: Optional[date] = None, end: Optional[date] = None, branch_id: Optional[int] = None, limit: int = 300):
    _guard(request, "finance.view")
    s, e = _period(start, end)
    qs = Expense.objects.filter(date__gte=s, date__lte=e).select_related("category", "branch", "created_by")
    if branch_id:
        qs = qs.filter(branch_id=branch_id)
    return qs[:limit]


@router.post("/expenses", response=ExpenseOut, auth=auth)
def create_expense(request, data: ExpenseIn):
    _guard(request, "finance.edit")
    if data.amount <= 0:
        raise HttpError(400, "Summa musbat bo'lsin.")
    x = Expense.objects.create(date=data.date or timezone.localdate(), branch_id=data.branch_id, category_id=data.category_id,
                               amount=data.amount, note=data.note, created_by=request.auth)
    record(request, "create", x)
    return x


@router.put("/expenses/{xid}", response=ExpenseOut, auth=auth)
def update_expense(request, xid: int, data: ExpenseIn):
    _guard(request, "finance.edit")
    x = get_object_or_404(Expense, pk=xid)
    for k, v in data.dict().items():
        if k == "date" and v is None:
            continue
        setattr(x, k, v)
    x.save()
    return x


@router.delete("/expenses/{xid}", auth=auth)
def delete_expense(request, xid: int):
    _guard(request, "finance.edit")
    x = get_object_or_404(Expense, pk=xid)
    record(request, "delete", x)
    x.delete()
    return {"ok": True}


@router.post("/expenses/{xid}/receipt", response=ExpenseOut, auth=auth)
def upload_receipt(request, xid: int, file: UploadedFile = File(...)):
    _guard(request, "finance.edit")
    x = get_object_or_404(Expense, pk=xid)
    x.receipt.save(file.name, file, save=True)
    return x


# ------------------------------------------------------------------ valyuta (Markaziy bank — real API)
@router.get("/rates", auth=auth)
def rates(request):
    """O'zbekiston Markaziy banki kurslari (cbu.uz) — import xomashyo narxini so'mga o'girish uchun."""
    _guard(request, "finance.view")
    from integrations.currency import get_rates

    return get_rates()
