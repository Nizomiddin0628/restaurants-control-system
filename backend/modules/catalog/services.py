"""Taomnoma xizmatlari: e'lon qilish (snapshot), tartiblash, Excel import."""
from __future__ import annotations

from io import BytesIO

from django.db import transaction
from django.db.models import Max

from core.events import emit

from .models import Category, MenuVersion, Modifier, ModifierGroup, Product


def build_snapshot() -> dict:
    """Faol va stop-listda bo'lmagan taomlar — sayt/Mini App/kassa uchun yagona JSON."""
    cats = []
    for c in Category.objects.filter(is_active=True, deleted_at__isnull=True):
        products = []
        for p in c.products.filter(is_active=True, deleted_at__isnull=True, in_stop_list=False).prefetch_related("modifier_groups__options", "branch_prices"):
            products.append({
                "id": p.pk, "name": p.name, "description": p.description, "price": p.price,
                "image": p.image.url if p.image else None, "weight_g": p.weight_g, "kcal": p.kcal, "tags": p.tags,
                "branch_prices": {bp.branch_id: bp.price for bp in p.branch_prices.all()},
                "modifier_groups": [
                    {"id": g.pk, "name": g.name, "min": g.min_select, "max": g.max_select,
                     "options": [{"id": m.pk, "name": m.name, "price": m.price, "default": m.is_default} for m in g.options.filter(is_active=True)]}
                    for g in p.modifier_groups.filter(is_active=True)
                ],
                "custom_data": p.custom_data,
            })
        if products:
            cats.append({"id": c.pk, "name": c.name, "image": c.image.url if c.image else None, "products": products})
    return {"categories": cats}


@transaction.atomic
def publish(by: str = "", note: str = "", tenant=None) -> MenuVersion:
    v = (MenuVersion.objects.aggregate(m=Max("version"))["m"] or 0) + 1
    mv = MenuVersion.objects.create(version=v, published_by=by, note=note, snapshot=build_snapshot())
    emit("catalog.published", {"version": v}, tenant=tenant)
    return mv


def current_snapshot() -> dict | None:
    mv = MenuVersion.objects.order_by("-version").first()
    return mv.snapshot if mv else None


@transaction.atomic
def reorder(model, ids: list[int]):
    for i, pk in enumerate(ids):
        model.objects.filter(pk=pk).update(sort_order=i)


def import_excel(file_bytes: bytes) -> dict:
    """
    Ustunlar (1-qator sarlavha, nomi bo'yicha, tartib muhim emas):
      kategoriya | nom_uz | nom_ru | nom_en | narx | tannarx | tavsif_uz | og'irlik | kkal | sku
    Kategoriya yo'q bo'lsa yaratiladi. Nom (uz) + kategoriya bo'yicha mavjud taom yangilanadi.
    """
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"created": 0, "updated": 0, "errors": ["Fayl bo'sh"]}
    header = [str(h or "").strip().lower() for h in rows[0]]
    col = {name: header.index(name) for name in header if name}
    need = ["kategoriya", "nom_uz", "narx"]
    missing = [n for n in need if n not in col]
    if missing:
        return {"created": 0, "updated": 0, "errors": [f"Ustunlar yetishmaydi: {', '.join(missing)}"]}

    created = updated = 0
    errors: list[str] = []
    cat_cache: dict[str, Category] = {}
    with transaction.atomic():
        for n, row in enumerate(rows[1:], start=2):
            try:
                def cell(name, default=None, row=row):
                    idx = col.get(name)
                    return row[idx] if idx is not None and idx < len(row) and row[idx] is not None else default

                cat_name = str(cell("kategoriya", "")).strip()
                name_uz = str(cell("nom_uz", "")).strip()
                if not cat_name or not name_uz:
                    continue
                cat = cat_cache.get(cat_name)
                if cat is None:
                    cat = Category.objects.filter(name__uz=cat_name, deleted_at__isnull=True).first()
                    if cat is None:
                        cat = Category.objects.create(name={"uz": cat_name, "ru": cat_name, "en": cat_name}, sort_order=Category.objects.count())
                    cat_cache[cat_name] = cat
                defaults = {
                    "name": {"uz": name_uz, "ru": str(cell("nom_ru", name_uz)), "en": str(cell("nom_en", name_uz))},
                    "price": int(float(cell("narx", 0))),
                    "cost": int(float(cell("tannarx", 0) or 0)),
                    "description": {"uz": str(cell("tavsif_uz", "") or ""), "ru": "", "en": ""},
                    "weight_g": int(cell("og'irlik")) if cell("og'irlik") else None,
                    "kcal": int(cell("kkal")) if cell("kkal") else None,
                    "sku": str(cell("sku", "") or ""),
                }
                p = Product.objects.filter(category=cat, name__uz=name_uz, deleted_at__isnull=True).first()
                if p:
                    for k, v in defaults.items():
                        setattr(p, k, v)
                    p.save()
                    updated += 1
                else:
                    Product.objects.create(category=cat, sort_order=cat.products.count(), **defaults)
                    created += 1
            except Exception as e:  # qator xatosi importni to'xtatmaydi, hisobotga tushadi
                errors.append(f"{n}-qator: {e}")
    return {"created": created, "updated": updated, "errors": errors}


__all__ = ["build_snapshot", "publish", "current_snapshot", "reorder", "import_excel", "Modifier", "ModifierGroup"]
