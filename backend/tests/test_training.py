"""
O'qitish va komplayens: auditoriya bo'yicha biriktirish, video o'tkazib yuborilmaydi, test bali,
kurs tugashi → sertifikat, topshiriq → dalil → tasdiq, standart tanishish, hisobot, modul o'chsa 404.
"""
from datetime import timedelta

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from django_tenants.utils import schema_context

H = {"HTTP_HOST": "lazzat.testserver"}


@pytest.fixture
def tr(tenant):
    from public.services import set_modules

    with schema_context("public"):
        set_modules(tenant, [*tenant.enabled_modules, "training"])
    with schema_context("lazzat"):
        from core.auth import issue_token
        from core.models import Membership, Role, User
        from modules.training.models import Course, Lesson, Question, Quiz

        cook_role, _ = Role.objects.get_or_create(code="cook", defaults={"name": "Oshpaz", "permissions": ["training.view"]})
        if "training.view" not in cook_role.permissions:
            cook_role.permissions = [*cook_role.permissions, "training.view"]
            cook_role.save()
        cook, _ = User.objects.get_or_create(phone="+998900000111", defaults={"full_name": "Azizbek Oshpaz"})
        Membership.objects.get_or_create(user=cook, role=cook_role)
        waiter_role, _ = Role.objects.get_or_create(code="waiter", defaults={"name": "Ofitsiant", "permissions": ["training.view"]})
        waiter, _ = User.objects.get_or_create(phone="+998900000222", defaults={"full_name": "Ofitsiant"})
        Membership.objects.get_or_create(user=waiter, role=waiter_role)

        Course.objects.filter(title__startswith="T:").delete()
        c = Course.objects.create(title="T: Burger", roles=["cook"], is_published=True, due_days=3, pass_score=80)
        l1 = Lesson.objects.create(course=c, title="Video dars", video_url="https://youtu.be/dQw4w9WgXcQ", sort_order=0)
        l2 = Lesson.objects.create(course=c, title="Matn dars", body="...", sort_order=1)
        q = Quiz.objects.create(course=c, title="Yakuniy", shuffle=False)
        q1 = Question.objects.create(quiz=q, text="Harorat?", options=[{"id": "a", "text": "55"}, {"id": "b", "text": "72"}], correct=["b"])
        q2 = Question.objects.create(quiz=q, text="Sous?", options=[{"id": "a", "text": "20 g"}, {"id": "b", "text": "50 g"}], correct=["a"])
        token = issue_token(cook, "lazzat")
        yield {"course": c, "l1": l1, "l2": l2, "quiz": q, "q1": q1, "q2": q2, "cook": cook, "waiter": waiter,
               "h": {**H, "HTTP_AUTHORIZATION": f"Bearer {token}"}}


def _post(client, url, data, h):
    return client.post(url, data=data, content_type="application/json", **h)


@pytest.mark.django_db
def test_audience_assigns_only_matching_roles(client, tr):
    r = client.get("/api/v1/training/my", **tr["h"])
    assert r.status_code == 200, r.content
    titles = [c["course"]["title"] for c in r.json()["courses"]]
    assert "T: Burger" in titles
    with schema_context("lazzat"):
        from modules.training.models import Enrollment
        assert Enrollment.objects.filter(course=tr["course"], user=tr["cook"]).exists()
        assert not Enrollment.objects.filter(course=tr["course"], user=tr["waiter"]).exists()


@pytest.mark.django_db
def test_video_cannot_be_skipped_and_lessons_are_sequential(client, tr):
    client.get("/api/v1/training/my", **tr["h"])
    # 2-dars 1-dars tugamaguncha yopiq
    assert client.get(f"/api/v1/training/my/lessons/{tr['l2'].pk}", **tr["h"]).status_code == 403
    assert client.get(f"/api/v1/training/my/lessons/{tr['l1'].pk}", **tr["h"]).status_code == 200
    url = f"/api/v1/training/my/lessons/{tr['l1'].pk}/beat"
    # oxiriga sakrash: 600 soniyalik videoda 590-soniyaga surdi, 5 soniya o'ynadi → foiz o'smaydi
    r = _post(client, url, {"position": 590, "duration": 600, "played": 5}, tr["h"])
    assert r.status_code == 200 and r.json()["percent"] < 5 and not r.json()["done"]
    # "Tugatdim" tugmasi video ko'rilmaguncha ishlamaydi
    assert _post(client, f"/api/v1/training/my/lessons/{tr['l1'].pk}/complete", {}, tr["h"]).status_code == 400
    # haqiqatda ko'rish: soat bo'yicha vaqt o'tdi deb hisoblaymiz
    with schema_context("lazzat"):
        from modules.training.models import LessonProgress
        from modules.training.services import beat
        for i in range(1, 61):
            LessonProgress.objects.filter(user=tr["cook"], lesson=tr["l1"]).update(last_beat_at=timezone.now() - timedelta(seconds=10))
            p = beat(tr["cook"], tr["l1"], position=i * 10, duration=600, played=10)
        assert p.percent >= 90 and p.completed_at is not None
    assert client.get(f"/api/v1/training/my/lessons/{tr['l2'].pk}", **tr["h"]).status_code == 200
    assert _post(client, f"/api/v1/training/my/lessons/{tr['l2'].pk}/complete", {}, tr["h"]).status_code == 200


@pytest.mark.django_db
def test_quiz_scoring_completion_and_certificate(client, tr):
    client.get("/api/v1/training/my", **tr["h"])
    with schema_context("lazzat"):
        from modules.training.services import complete_text_lesson
        for les in (tr["l1"], tr["l2"]):
            complete_text_lesson(tr["cook"], les)
    r = _post(client, f"/api/v1/training/my/quizzes/{tr['quiz'].pk}/start", {}, tr["h"])
    assert r.status_code == 200, r.content
    assert "correct" not in str(r.json()["questions"])            # javoblar xodimga ko'rinmaydi
    aid = r.json()["attempt_id"]
    # 1 ta to'g'ri, 1 ta xato → 50% → o'tmadi
    r = _post(client, f"/api/v1/training/my/attempts/{aid}/submit", {"answers": {str(tr["q1"].pk): ["b"], str(tr["q2"].pk): ["b"]}}, tr["h"])
    assert r.json()["score"] == 50 and not r.json()["passed"] and not r.json()["course_completed"]
    aid = _post(client, f"/api/v1/training/my/quizzes/{tr['quiz'].pk}/start", {}, tr["h"]).json()["attempt_id"]
    r = _post(client, f"/api/v1/training/my/attempts/{aid}/submit", {"answers": {str(tr["q1"].pk): ["b"], str(tr["q2"].pk): ["a"]}}, tr["h"])
    assert r.json()["score"] == 100 and r.json()["passed"] and r.json()["course_completed"]
    cert = client.get(f"/api/v1/training/my/certificates/{r.json()['enrollment_id']}", **tr["h"])
    assert cert.status_code == 200 and cert.json()["user"] == "Azizbek Oshpaz" and cert.json()["score"] == 100


@pytest.mark.django_db
def test_assignment_proof_and_review(client, api, tr):
    r = api.post("/api/v1/training/assignments", {"title": "Burger rasmi", "roles": ["cook"], "requires_proof": True})
    assert r.status_code == 200 and r.json()["assigned"] >= 1, r.content
    subs = client.get("/api/v1/training/my/assignments", **tr["h"]).json()
    sid = next(s["id"] for s in subs if s["assignment"]["title"] == "Burger rasmi")
    # dalilsiz topshirib bo'lmaydi
    assert _post(client, f"/api/v1/training/my/assignments/{sid}/submit", {"text": "tayyor"}, tr["h"]).status_code == 400
    img = SimpleUploadedFile("burger.jpg", b"\xff\xd8\xff\xe0fakejpeg", content_type="image/jpeg")
    r = client.post(f"/api/v1/training/my/assignments/{sid}/files", {"file": img}, **tr["h"])
    assert r.status_code == 200 and len(r.json()["files"]) == 1, r.content
    assert _post(client, f"/api/v1/training/my/assignments/{sid}/submit", {"text": "tayyor"}, tr["h"]).json()["status"] == "submitted"
    # sababsiz qaytarib bo'lmaydi
    assert api.post(f"/api/v1/training/submissions/{sid}/review", {"approve": False}).status_code == 400
    r = api.post(f"/api/v1/training/submissions/{sid}/review", {"approve": True, "note": "Zo'r"})
    assert r.status_code == 200 and r.json()["status"] == "approved"
    # oddiy xodim boshqalarni tekshira olmaydi
    aid = r.json()["assignment"]["id"]
    assert client.get(f"/api/v1/training/assignments/{aid}/submissions", **tr["h"]).status_code == 403


@pytest.mark.django_db
def test_standard_ack_and_new_version(client, api, tr):
    r = api.post("/api/v1/training/standards", {"title": "T: Qo'les yuvish", "everyone": True, "body": "20 soniya"})
    sid = r.json()["id"]
    mine = client.get("/api/v1/training/my/standards", **tr["h"]).json()
    assert any(s["id"] == sid and s["acked_at"] is None for s in mine)
    assert _post(client, f"/api/v1/training/my/standards/{sid}/ack", {}, tr["h"]).json()["acked_at"]
    api.post(f"/api/v1/training/standards/{sid}/new-version")
    mine = client.get("/api/v1/training/my/standards", **tr["h"]).json()
    assert next(s for s in mine if s["id"] == sid)["acked_at"] is None     # yangi versiya — qayta tanishish


@pytest.mark.django_db
def test_report_and_permissions(client, api, tr):
    client.get("/api/v1/training/my", **tr["h"])
    r = api.get("/api/v1/training/report/summary")
    assert r.status_code == 200
    row = next(x for x in r.json()["rows"] if x["user"]["phone"] == "+998900000111")
    assert row["courses_total"] >= 1 and row["lessons_total"] >= 2
    assert api.get(f"/api/v1/training/report/users/{tr['cook'].pk}").status_code == 200
    # xodim hisobotni va boshqaruvni ko'ra olmaydi
    assert client.get("/api/v1/training/report/summary", **tr["h"]).status_code == 403
    assert _post(client, "/api/v1/training/courses", {"title": "x"}, tr["h"]).status_code == 403


@pytest.mark.django_db
def test_admin_builds_course_and_quiz(api, tr):
    c = api.post("/api/v1/training/courses", {"title": "T: Yangi", "roles": ["cook"], "is_published": True}).json()
    les = api.post(f"/api/v1/training/courses/{c['id']}/lessons", {"title": "1-dars", "checklist": ["a", "b"]}).json()
    assert les["checklist"] == ["a", "b"]
    q = api.post(f"/api/v1/training/courses/{c['id']}/quizzes", {"title": "Test"}).json()
    bad = api.put(f"/api/v1/training/quizzes/{q['id']}/questions", {"questions": [{"text": "?", "options": [{"id": "a", "text": "1"}], "correct": ["a"]}]})
    assert bad.status_code == 400
    ok = api.put(f"/api/v1/training/quizzes/{q['id']}/questions", {"questions": [
        {"text": "2+2?", "options": [{"id": "a", "text": "3"}, {"id": "b", "text": "4"}], "correct": ["b"]}]})
    assert ok.status_code == 200 and ok.json()["questions_count"] == 1
    full = api.get(f"/api/v1/training/courses/{c['id']}").json()
    assert full["enrolled"] >= 1 and len(full["lessons"]) == 1


@pytest.mark.django_db
def test_module_disabled_returns_404(api, tenant, tr):
    from public.services import set_modules

    with schema_context("public"):
        set_modules(tenant, [m for m in tenant.enabled_modules if m != "training"])
    assert api.get("/api/v1/training/my").status_code == 404
    with schema_context("public"):
        set_modules(tenant, [*tenant.enabled_modules, "training"])


def test_media_link_detection():
    from modules.training.media import info
    assert info("https://youtu.be/dQw4w9WgXcQ")["kind"] == "youtube"
    assert info("https://www.youtube.com/watch?v=dQw4w9WgXcQ&t=5")["id"] == "dQw4w9WgXcQ"
    d = info("https://drive.google.com/file/d/1AbCdEfGhIjKlMnOpQrStUvWxYz012345/view?usp=sharing")
    assert d["kind"] == "drive" and d["embed"].endswith("/preview")
    assert info("https://vimeo.com/123456789")["kind"] == "vimeo"
    assert info("https://cdn.site.uz/burger.MP4?x=1")["kind"] == "video"
    assert info("https://site.uz/a.jpg")["kind"] == "image"
    assert info("https://site.uz/page")["kind"] == "link"


@pytest.mark.django_db
def test_links_instead_of_uploads(client, api, tr):
    """Material havola bilan: dars rasmi, fayl-havola, Drive video (vaqt bo'yicha nazorat), kurs muqovasi."""
    r = api.put(f"/api/v1/training/lessons/{tr['l2'].pk}", {"title": "Matn dars", "image_url": "https://site.uz/burger.png",
                                                              "video_url": "https://drive.google.com/file/d/1AbCdEfGhIjKlMnOpQrStUvWxYz012345/view",
                                                              "duration_seconds": 120})
    assert r.status_code == 200 and r.json()["video_mode"] == "time" and r.json()["image"] == "https://site.uz/burger.png"
    assert api.post(f"/api/v1/training/lessons/{tr['l2'].pk}/links", {"url": "ftp://x"}).status_code == 400
    r = api.post(f"/api/v1/training/lessons/{tr['l2'].pk}/links", {"url": "https://site.uz/qollanma.pdf", "title": "Qo'llanma"})
    assert r.status_code == 200 and r.json()["files"][-1]["is_link"] and r.json()["files"][-1]["title"] == "Qo'llanma"
    c = api.put(f"/api/v1/training/courses/{tr['course'].pk}", {"title": "T: Burger", "roles": ["cook"], "is_published": True,
                                                                 "cover_url": "https://site.uz/c.jpg"})
    assert c.json()["cover"] == "https://site.uz/c.jpg"
    # Drive videoda: brauzer yuborgan davomiylik emas, admin yozgan 120 s ishlatiladi
    with schema_context("lazzat"):
        from modules.training.models import Lesson
        from modules.training.services import beat
        les = Lesson.objects.get(pk=tr["l2"].pk)
        p = beat(tr["cook"], les, position=10, duration=5, played=10)
        assert p.duration == 120 and not p.completed_at


@pytest.mark.django_db
def test_quiz_locked_until_lessons_done_and_resumes(client, tr):
    """Darslar tugamaguncha test ochilmaydi; sahifa yangilansa — o'sha urinish davom etadi (taymer qaytadan boshlanmaydi),
    javoblar saqlanadi; vaqti o'tgan urinish urinish sifatida hisoblanadi."""
    client.get("/api/v1/training/my", **tr["h"])
    url = f"/api/v1/training/my/quizzes/{tr['quiz'].pk}/start"
    assert _post(client, url, {}, tr["h"]).status_code == 403
    assert _post(client, f"/api/v1/training/my/lessons/{tr['l2'].pk}/complete", {}, tr["h"]).status_code == 403   # ketma-ketlik
    with schema_context("lazzat"):
        from modules.training.models import Quiz
        from modules.training.services import complete_text_lesson
        for les in (tr["l1"], tr["l2"]):
            complete_text_lesson(tr["cook"], les)
        Quiz.objects.filter(pk=tr["quiz"].pk).update(time_limit_seconds=120, max_attempts=2)
    a1 = _post(client, url, {}, tr["h"]).json()
    assert a1["remaining_seconds"] and a1["remaining_seconds"] <= 120
    _post(client, f"/api/v1/training/my/attempts/{a1['attempt_id']}/save", {"answers": {str(tr["q1"].pk): ["b"]}}, tr["h"])
    a2 = _post(client, url, {}, tr["h"]).json()                                   # sahifa yangilandi
    assert a2["attempt_id"] == a1["attempt_id"] and a2["answers"] == {str(tr["q1"].pk): ["b"]}
    with schema_context("lazzat"):
        from modules.training.models import QuizAttempt
        QuizAttempt.objects.filter(pk=a1["attempt_id"]).update(started_at=timezone.now() - timedelta(minutes=10))
    a3 = _post(client, url, {}, tr["h"]).json()                                   # eski urinish vaqti o'tdi → hisoblandi
    assert a3["attempt_id"] != a1["attempt_id"] and a3["attempts_left"] == 1
    with schema_context("lazzat"):
        old = QuizAttempt.objects.get(pk=a1["attempt_id"])
        assert old.finished_at is not None and old.score == 0                      # kech topshirilgan — ball berilmaydi


@pytest.mark.django_db
def test_reminders_to_employee_and_responsible(client, api, tr, tenant, monkeypatch):
    sent = []
    monkeypatch.setattr("modules.training.services.notify", lambda t, u, text: sent.append((u.phone, text)))
    with schema_context("lazzat"):
        from core.models import User
        from modules.training.models import Course, Enrollment, Reminder
        boss = User.objects.get(phone="+998901234567")
        User.objects.filter(pk=tr["cook"].pk).update(telegram_id=111)
        Course.objects.filter(pk=tr["course"].pk).update(responsible=boss)
        Reminder.objects.all().delete()
    client.get("/api/v1/training/my", **tr["h"])
    with schema_context("lazzat"):
        Enrollment.objects.filter(course=tr["course"], user=tr["cook"]).update(due_at=timezone.now() - timedelta(days=1))
    r = api.post("/api/v1/training/report/remind").json()
    assert r["users"] >= 1 and r["responsibles"] >= 1
    assert any(p == "+998900000111" and "muddati o'tdi" in t for p, t in sent)
    assert any(p == "+998901234567" and "Azizbek Oshpaz" in t for p, t in sent)
    n = len(sent)
    api.post("/api/v1/training/report/remind")                                     # bir kunda ikkinchi marta — spam yo'q
    assert len(sent) == n


@pytest.mark.django_db
def test_report_export_xlsx(api, tr):
    r = api.get("/api/v1/training/report/export.xlsx")
    assert r.status_code == 200 and r["Content-Type"].startswith("application/vnd.openxmlformats")
    from io import BytesIO

    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Xodimlar", "Kurslar bo'yicha"] and wb["Xodimlar"]["A1"].value == "Xodim"
